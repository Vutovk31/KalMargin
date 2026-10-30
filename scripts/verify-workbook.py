"""Independent acceptance oracle; never prints commercial values or source filename."""
import sys
from io import BytesIO
from pathlib import Path
from decimal import Decimal, ROUND_CEILING, ROUND_FLOOR, getcontext
import openpyxl

getcontext().prec = 60
source_path, export_path, target_text = sys.argv[1:]
target = Decimal(target_text)
source = openpyxl.load_workbook(source_path, data_only=True)
expected = []
summary_before = summary_after = summary_cost = Decimal(0)
for sheet in source:
    for cells in sheet.iter_rows():
        values = [c.value for c in cells]
        names = ['Код товара', 'Кол-во', 'Цена со скидкой', 'Закупочная цена товара']
        if all(name in values for name in names):
            columns = {name: values.index(name) for name in names}
            for row in sheet.iter_rows(min_row=cells[0].row + 1, values_only=True):
                if not row[columns['Код товара']]:
                    continue
                price = Decimal(str(row[columns['Цена со скидкой']]))
                cost = Decimal(str(row[columns['Закупочная цена товара']]))
                qty = Decimal(str(row[columns['Кол-во']]))
                cap = int(price.to_integral_value(rounding=ROUND_FLOOR))
                assert cap >= 1, 'Source needs an impossible positive whole-ruble price'
                minimum = max(1, int((cost / (1-target/100)).to_integral_value(rounding=ROUND_CEILING)))
                final = min(cap, minimum)
                expected.append(final)
                summary_before += price * qty
                summary_after += final * qty
                summary_cost += cost * qty
            break
    if expected:
        break
assert expected, 'No source rows found'
output = openpyxl.load_workbook(BytesIO(Path(export_path).read_bytes()), data_only=False)
result = output['Результат']
assert result.max_row == len(expected) + 1, 'Output row count mismatch'
for index, final in enumerate(expected, 2):
    cell = result.cell(index, 2)
    assert cell.data_type == 'n', 'Output price is not numeric'
    assert isinstance(cell.value, int) and cell.value >= 1, 'Fractional or nonpositive output price'
    assert cell.value == final, 'Independent calculation mismatch'
summary = dict(output['Сводка'].iter_rows(min_row=2, values_only=True))
for key, value in [('before', summary_before), ('sum', summary_after), ('cost', summary_cost), ('profit', summary_after-summary_cost), ('beforeProfit', summary_before-summary_cost)]:
    assert Decimal(str(summary[key])) == value, 'Exact summary mismatch: ' + key
print('Independent Decimal/openpyxl verification passed; rows=' + str(len(expected)))
